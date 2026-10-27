"""Loaders for prose documents: PDF, HTML, Markdown/TXT and DOCX."""

from __future__ import annotations

import io
import re
from typing import Any

import yaml

from app.core.text import normalize_whitespace
from app.ingestion.loaders.base import (
    Block,
    LoadedDocument,
    LoaderError,
    decode_text,
    is_backmatter_heading,
    is_reference_heading,
)

_NUMBERED_HEADING = re.compile(r"^(?:\d+(?:\.\d+)*\.?|[IVX]+\.)\s+[A-ZÀ-Ý][^.!?]{2,80}$")
_KNOWN_SECTIONS = re.compile(
    r"^(abstract|summary|background|introduction|methods?|materials and methods|results|discussion|"
    r"conclusions?|key points|recommendations?|résumé|introduction|méthodes|résultats|conclusion|"
    r"resumen|introducción|métodos|resultados|discusión|conclusiones|resumo|introdução|discussão|"
    r"conclusões|zusammenfassung|einleitung|methoden|ergebnisse|diskussion|schlussfolgerung(en)?)\s*:?$",
    re.IGNORECASE,
)


def _looks_like_heading(line: str) -> bool:
    stripped = line.strip()
    if not stripped or len(stripped) > 90 or stripped.endswith((",", ";")):
        return False
    if _KNOWN_SECTIONS.match(stripped) or _NUMBERED_HEADING.match(stripped):
        return True
    letters = [c for c in stripped if c.isalpha()]
    return len(letters) >= 4 and len(stripped.split()) <= 8 and all(c.isupper() for c in letters)


# ── PDF ─────────────────────────────────────────────────────────────────────
def load_pdf(data: bytes, path: str) -> list[LoadedDocument]:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise LoaderError("encrypted PDF") from exc
    except LoaderError:
        raise
    except Exception as exc:
        raise LoaderError(f"unreadable PDF: {exc}") from exc

    meta: dict[str, Any] = {}
    info = reader.metadata or {}
    title = str(info.get("/Title") or "").strip()
    if title and not re.search(r"(microsoft word|untitled|\.docx?$|\.pdf$)", title, re.I):
        meta["title"] = title
    if info.get("/Author"):
        meta["authors"] = [a.strip() for a in re.split(r";|,(?=\s*[A-Z])| and ", str(info["/Author"])) if a.strip()]
    created = str(info.get("/CreationDate") or "")
    if m := re.search(r"D:((?:19|20)\d\d)", created):
        meta["_pdf_year"] = int(m.group(1))

    blocks: list[Block] = []
    section: str | None = None
    total_pages = len(reader.pages)
    in_references = False
    for page_no, page in enumerate(reader.pages, start=1):
        try:
            raw = page.extract_text() or ""
        except Exception:
            continue
        lines, buffer = normalize_whitespace(raw).split("\n"), []
        for line in lines:
            if is_reference_heading(line) and page_no >= max(2, total_pages * 0.4):
                in_references = True  # drop reference lists: they pollute retrieval with off-topic titles
                break
            if is_backmatter_heading(line) and page_no >= total_pages * 0.5:
                in_references = True
                break
            if _looks_like_heading(line):
                if buffer:
                    blocks.append(Block(" ".join(buffer), section, page_no))
                    buffer = []
                section = line.strip().rstrip(":").title() if line.isupper() else line.strip().rstrip(":")
                continue
            if not line.strip():
                if buffer:
                    blocks.append(Block(" ".join(buffer), section, page_no))
                    buffer = []
                continue
            buffer.append(line.strip())
        if buffer:
            blocks.append(Block(" ".join(buffer), section, page_no))
        if in_references:
            break
    if not any(b.text.strip() for b in blocks):
        raise LoaderError("no extractable text (scanned PDF? run OCR first, e.g. `ocrmypdf in.pdf out.pdf`)")
    meta["_pages"] = total_pages
    return [LoadedDocument(blocks=_merge_small_blocks(blocks), metadata=meta)]


def _merge_small_blocks(blocks: list[Block], min_chars: int = 200) -> list[Block]:
    """PDF extraction yields many line-level blocks; merge consecutive ones in the same section/page."""
    merged: list[Block] = []
    for block in blocks:
        prev = merged[-1] if merged else None
        if prev and prev.section == block.section and prev.page == block.page and len(prev.text) < min_chars:
            prev.text = f"{prev.text} {block.text}"
        else:
            merged.append(Block(block.text, block.section, block.page))
    return merged


# ── HTML ────────────────────────────────────────────────────────────────────
_BOILERPLATE = re.compile(
    r"(cookie|consent|banner|newsletter|share|social|breadcrumb|menu|navbar|sidebar|footer|header|related|"
    r"advert|promo|subscribe|comment|popup|modal|skip-link)",
    re.IGNORECASE,
)


def load_html(data: bytes, path: str) -> list[LoadedDocument]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(decode_text(data), "lxml")
    meta: dict[str, Any] = {}

    def meta_values(*names: str) -> list[str]:
        values = []
        for name in names:
            for tag in soup.find_all("meta", attrs={"name": name}) + soup.find_all("meta", attrs={"property": name}):
                content = (tag.get("content") or "").strip()
                if content:
                    values.append(content)
        return values

    if v := meta_values("citation_title", "dc.title", "DC.title", "og:title"):
        meta["title"] = v[0]
    elif soup.title and soup.title.string:
        meta["title"] = soup.title.string.strip()
    if v := meta_values("citation_author", "dc.creator", "DC.creator", "author"):
        meta["authors"] = v
    if v := meta_values("citation_publication_date", "citation_date", "dc.date", "DC.date",
                        "article:published_time", "datePublished"):
        meta["published_date"] = v[0][:10]
    if v := meta_values("citation_journal_title"):
        meta["journal"] = v[0]
    if v := meta_values("citation_doi", "dc.identifier"):
        meta["doi"] = v[0].removeprefix("doi:")
    if v := meta_values("citation_pmid"):
        meta["pmid"] = v[0]
    if v := meta_values("og:site_name", "citation_publisher", "dc.publisher"):
        meta["publisher"] = v[0]
    if v := meta_values("og:url"):
        meta["url"] = v[0]
    if canonical := soup.find("link", rel="canonical"):
        meta["url"] = canonical.get("href") or meta.get("url")
    if soup.html and soup.html.get("lang"):
        meta["language"] = soup.html["lang"][:2].lower()

    for tag in soup(["script", "style", "noscript", "svg", "nav", "footer", "header", "aside", "form", "iframe",
                     "button", "figure"]):
        tag.decompose()
    for tag in soup.find_all(True):
        if tag.attrs is None:
            continue
        marker = " ".join(tag.get("class", [])) + " " + (tag.get("id") or "") + " " + (tag.get("role") or "")
        if marker.strip() and _BOILERPLATE.search(marker) and tag.name not in ("main", "article", "body", "html"):
            tag.decompose()

    root = soup.find("article") or soup.find("main") or soup.find(attrs={"role": "main"}) or soup.body or soup
    blocks: list[Block] = []
    section: str | None = None
    for el in root.find_all(["h1", "h2", "h3", "h4", "p", "li", "blockquote", "td", "dd", "pre"]):
        text = normalize_whitespace(el.get_text(" ", strip=True))
        if not text:
            continue
        if el.name in ("h1", "h2", "h3", "h4"):
            if is_reference_heading(text):
                break
            section = text[:200]
            continue
        if el.find_parent(["li", "td", "blockquote"]) and el.name == "p":
            continue  # avoid duplicating nested content
        blocks.append(Block(text, section))
    if not blocks:
        text = normalize_whitespace(root.get_text("\n", strip=True))
        if text:
            blocks = [Block(text)]
    if not blocks:
        raise LoaderError("no textual content in HTML")
    return [LoadedDocument(blocks=_merge_small_blocks(blocks, 400), metadata=meta)]


# ── Markdown / plain text ───────────────────────────────────────────────────
_FRONT_MATTER = re.compile(r"\A---\s*\n(.*?)\n---\s*\n", re.DOTALL)
_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_MD_INLINE = [
    (re.compile(r"!\[[^\]]*\]\([^)]*\)"), ""),
    (re.compile(r"\[([^\]]+)\]\([^)]*\)"), r"\1"),
    (re.compile(r"(\*\*|__)(.+?)\1"), r"\2"),
    (re.compile(r"(?<!\w)[*_](.+?)[*_](?!\w)"), r"\1"),
    (re.compile(r"`([^`]+)`"), r"\1"),
    (re.compile(r"^\s*[-*+]\s+", re.M), "• "),
    (re.compile(r"^\s*>\s?", re.M), ""),
]


def load_text(data: bytes, path: str) -> list[LoadedDocument]:
    text = decode_text(data)
    meta: dict[str, Any] = {}
    if match := _FRONT_MATTER.match(text):
        try:
            front = yaml.safe_load(match.group(1)) or {}
            if isinstance(front, dict):
                meta.update(front)
        except yaml.YAMLError:
            pass
        text = text[match.end():]
    is_markdown = path.lower().endswith((".md", ".markdown"))
    blocks: list[Block] = []
    section: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            body = "\n".join(buffer).strip()
            if body:
                blocks.append(Block(normalize_whitespace(body), section))
            buffer.clear()

    for line in text.splitlines():
        heading = _MD_HEADING.match(line) if is_markdown else None
        if heading:
            flush()
            title = heading.group(2).strip()
            if is_reference_heading(title):
                break
            if len(heading.group(1)) == 1 and "title" not in meta:
                meta["title"] = title
            section = title
        elif not line.strip():
            flush()
        else:
            if is_markdown:
                for pattern, repl in _MD_INLINE:
                    line = pattern.sub(repl, line)
            buffer.append(line)
    flush()
    if not blocks:
        raise LoaderError("empty text document")
    return [LoadedDocument(blocks=blocks, metadata=meta)]


# ── DOCX ────────────────────────────────────────────────────────────────────
def load_docx(data: bytes, path: str) -> list[LoadedDocument]:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise LoaderError("python-docx is not installed") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise LoaderError(f"unreadable DOCX: {exc}") from exc
    meta: dict[str, Any] = {}
    props = document.core_properties
    if props.title:
        meta["title"] = props.title
    if props.author:
        meta["authors"] = [a.strip() for a in re.split(r";|,", props.author) if a.strip()]
    if props.created:
        meta["_doc_year"] = props.created.year
    blocks: list[Block] = []
    section: str | None = None
    for para in document.paragraphs:
        text = normalize_whitespace(para.text)
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith(("heading", "title", "titre", "überschrift")):
            if is_reference_heading(text):
                break
            if style.startswith(("title", "titre")) and "title" not in meta:
                meta["title"] = text
            section = text
            continue
        blocks.append(Block(text, section))
    for table in document.tables:
        for row in table.rows:
            cells = [normalize_whitespace(c.text) for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(Block(" | ".join(dict.fromkeys(cells)), section))
    if not blocks:
        raise LoaderError("empty DOCX")
    return [LoadedDocument(blocks=_merge_small_blocks(blocks, 400), metadata=meta)]
